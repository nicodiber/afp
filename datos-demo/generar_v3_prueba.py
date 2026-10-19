"""Genera datos-demo/v3-prueba.xlsx: planilla FICTICIA con el formato de la hoja «movimientos» de AEP v3.

Mismo formato que v3 (dos filas de encabezado, columnas de categorías con emoji, tres columnas por
cuenta: Entrada, Salida y Saldo), con cuentas genéricas y los 21 casos de PLAN_DE_VERSIONES.md §7.
Uso: python3 datos-demo/generar_v3_prueba.py [cantidad_de_filas_extra]
Sin número, guarda también v3-prueba-respuestas.csv (3 respuestas nuevas del Form, PLAN v0.7.0-A8).
Con un número, agrega filas sintéticas al final (para medir velocidad) y guarda v3-prueba-grande.xlsx.
"""
import datetime as dt
import random
import sys
from pathlib import Path

import openpyxl

INGRESOS = ['💼 Trabajo', '👪 Aporte familiar', '🐛 Mantenido', '📈 Rendimiento', '📝 i_Préstamo',
            '🧾 i_Deuda', '🔙 i_Inversión', '🎁 i_Regalo']
EGRESOS = ['🛒 Supermercado', '🥦 Verdulería', '🍫 Kiosco', '🏠 Vivienda: Alquiler', '🏖️ Vacaciones',
           '📝 e_Préstamo', '🧾 e_Deuda', '🔜 e_Inversión', '🔧 Reparación']
CUENTAS = [('Billetera', '👛', 10_000), ('Banco', '🏛️', 100_000), ('Billetera virtual', '🤝', 50_000)]

# Encabezados
fila1, fila2 = [None, None, None], ['Marca temporal', 'Fecha', 'Descripción']
fila1 += ['Ingreso'] + [None] * (len(INGRESOS) - 1); fila2 += INGRESOS
fila1 += ['Egreso'] + [None] * (len(EGRESOS) - 1); fila2 += EGRESOS
for nombre, emoji, _ in CUENTAS:
    fila1 += [nombre, None, None]; fila2 += [f'{emoji}➕ {nombre} Entrada', f'{emoji}➖ {nombre} Salida', 'Saldo']
fila1 += ['Cuentas por cobrar', None, None, None]; fila2 += ['👤 Persona', '📤 Debe', '📥 Haber', 'Saldo']
fila1 += ['Cuentas por pagar', None, None, None]; fila2 += ['👤 Persona', '📥 Debe', '📤 Haber', 'Saldo']
fila1 += ['USD', None, None]; fila2 += ['📥 Debe', '📤 Haber', 'Saldo']
fila1 += [None, None]; fila2 += ['Ocultar (Email)', 'Acción']
col = {h: i for i, h in enumerate(fila2) if h != 'Saldo'}
def c_cuenta(nombre, lado):
    emoji = next(e for n, e, _ in CUENTAS if n == nombre)
    return col[f'{emoji}{"➕" if lado == "E" else "➖"} {nombre} {"Entrada" if lado == "E" else "Salida"}']

EGR, ING, TRA = '💸 Egreso', '💰 Ingreso', '🔄 Transacción entre las billeteras'
# (descripción, {columna: valor}, acción)
CASOS = [
    ('Sueldo', {'💼 Trabajo': 500000, ('Banco', 'E'): 500000}, ING),
    ('Ayuda', {'👪 Aporte familiar': 100000, ('Billetera virtual', 'E'): 100000}, ING),
    ('Compra semanal', {'🛒 Supermercado': 30000, ('Billetera virtual', 'S'): 30000}, EGR),
    ('Feria y kiosco', {'🥦 Verdulería': 5000, '🍫 Kiosco': 1000, ('Billetera', 'S'): 6000}, EGR),
    ('Alquiler octubre', {'🏠 Vivienda: Alquiler': 300000, ('Banco', 'S'): 300000}, EGR),
    ('Viaje pagado por terceros', {'🐛 Mantenido': 20000, '🏖️ Vacaciones': 20000}, EGR),
    ('Extracción', {('Banco', 'S'): 50000, ('Billetera', 'E'): 50000}, TRA),
    ('Compra de dólares', {('Banco', 'S'): 145000, 'USD:📥 Debe': 100}, TRA),
    ('Venta de dólares', {'USD:📤 Haber': 50, ('Billetera', 'E'): 70000}, TRA),
    ('Préstamo a Persona A', {'📝 e_Préstamo': 15000, ('Billetera', 'S'): 15000, 'CxC:👤 Persona': 'Persona A', 'CxC:📤 Debe': 15000}, EGR),
    ('Devolución de Persona A', {'📝 i_Préstamo': 15000, ('Billetera', 'E'): 15000, 'CxC:👤 Persona': 'Persona A', 'CxC:📥 Haber': 15000}, ING),
    ('Préstamo de Persona B', {'🧾 i_Deuda': 8000, ('Billetera virtual', 'E'): 8000, 'CxP:👤 Persona': 'Persona B', 'CxP:📤 Haber': 8000}, ING),
    ('Pago a Persona B', {'🧾 e_Deuda': 8000, ('Billetera virtual', 'S'): 8000, 'CxP:👤 Persona': 'Persona B', 'CxP:📥 Debe': 8000}, EGR),
    ('Suscripción a fondo', {'🔜 e_Inversión': 40000, ('Banco', 'S'): 40000}, EGR),
    ('Rescate de fondo', {'🔙 i_Inversión': 10000, ('Banco', 'E'): 10000}, ING),
    ('Rendimiento diario', {'📈 Rendimiento': 1200, ('Billetera virtual', 'E'): 1200}, ING),
    ('Compra cargada a mano', {'🛒 Supermercado': 12000, ('Billetera virtual', 'S'): 12000}, None),
    ('Kiosco sin cuenta', {'🍫 Kiosco': 2000}, EGR),
    ('Fila sin datos', {}, None),
    ('Regalo de cumpleaños', {'🎁 i_Regalo': 5000, ('Billetera', 'E'): 5000}, ING),
    ('Arreglo de canilla', {'🔧 Reparación': 4000, ('Billetera', 'S'): 4000}, EGR),
]
GRUPOS = {'CxC': fila2.index('👤 Persona'), 'CxP': len(fila2) - 1 - fila2[::-1].index('👤 Persona'),
          'USD': fila1.index('USD')}

def columna(clave):
    if isinstance(clave, tuple):
        return c_cuenta(*clave)
    if ':' in clave and clave.split(':')[0] in ('CxC', 'CxP', 'USD'):
        grupo, nombre = clave.split(':', 1)
        base = GRUPOS[grupo]
        return base + fila2[base:base + 4].index(nombre)
    return col[clave]

def armar(extra=0):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'movimientos'
    ws.append(fila1); ws.append(fila2)
    saldos = {n: s for n, _, s in CUENTAS}
    saldos.update({'USD': 200, 'CxC': 0, 'CxP': 0})
    casos = list(CASOS)
    rnd = random.Random(7)
    for _ in range(extra):
        m = rnd.choice([1000, 2500, 4000, 9000, 15000])
        casos.append(('Gasto sintético', {'🛒 Supermercado': m, ('Billetera virtual', 'S'): m}, EGR))
    inicio = dt.datetime(2026, 10, 1, 9, 0)
    for i, (desc, valores, accion) in enumerate(casos):
        fila = [None] * len(fila2)
        sin_fecha = desc == 'Fila sin datos'
        fila[0] = inicio + dt.timedelta(hours=6 * i)
        fila[1] = None if sin_fecha else (inicio + dt.timedelta(hours=6 * i)).date()
        fila[2] = desc
        for clave, v in valores.items():
            fila[columna(clave)] = v
        for nombre, emoji, _ in CUENTAS:
            e = fila[c_cuenta(nombre, 'E')] or 0; s = fila[c_cuenta(nombre, 'S')] or 0
            saldos[nombre] += e - s
            fila[c_cuenta(nombre, 'S') + 1] = saldos[nombre]
        for grupo, debe, haber in (('USD', 0, 1), ('CxC', 1, 2), ('CxP', 1, 2)):
            base = GRUPOS[grupo]
            saldos[grupo] += (fila[base + debe] or 0) - (fila[base + haber] or 0)
            fila[base + haber + 1] = saldos[grupo]
        fila[-1] = accion
        ws.append(fila)
    for celda in ws['A'][2:]: celda.number_format = 'm/d/yyyy h:mm:ss'
    for celda in ws['B'][2:]: celda.number_format = 'd/M/yyyy'
    return wb

RESPUESTAS = [
    ('Compra rápida', {'🛒 Supermercado': 6000, ('Billetera virtual', 'S'): 6000}, EGR),
    ('Golosinas', {'🍫 Kiosco': 1500, ('Billetera', 'S'): 1500}, EGR),
    ('Changa', {'💼 Trabajo': 20000, ('Billetera', 'E'): 20000}, ING),
]

def csv_respuestas(destino):
    """Respuestas nuevas del Form, como las exporta Google Sheets a CSV (texto, dos filas de encabezado)."""
    import csv
    with open(destino, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['' if v is None else v for v in fila1]); w.writerow(fila2)
        for i, (desc, valores, accion) in enumerate(RESPUESTAS):
            fila = [''] * len(fila2)
            momento = dt.datetime(2026, 10, 30, 10 + i, 15, 30)
            fila[0] = f'{momento.month}/{momento.day}/{momento.year} {momento.hour}:{momento.minute:02d}:{momento.second:02d}'
            fila[1] = f'{momento.day}/{momento.month}/{momento.year}'
            fila[2] = desc
            for clave, v in valores.items():
                fila[columna(clave)] = v
            fila[-1] = accion
            w.writerow(fila)

if __name__ == '__main__':
    extra = int(sys.argv[1]) if len(sys.argv) > 1 else 0
    destino = Path(__file__).with_name('v3-prueba-grande.xlsx' if extra else 'v3-prueba.xlsx')
    armar(extra).save(destino)
    if not extra:
        csv_respuestas(destino.with_name('v3-prueba-respuestas.csv'))
    print(destino, len(CASOS) + extra, 'filas')
